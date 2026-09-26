import io

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from src.word_document import UnsafeWordEdit, import_content, render_edited_word


UPLOAD_ID = "a" * 32 + ".docx"


def _sample_word():
    document = Document()
    document.add_heading("Quarterly report", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Revenue ")
    emphasized = paragraph.add_run("increased")
    emphasized.bold = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table value"
    document.sections[0].header.paragraphs[0].text = "Confidential"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_unchanged_word_export_is_exact_original_binary():
    original = _sample_word()
    content = import_content(original, UPLOAD_ID)
    assert render_edited_word(original, content, content) == original


def test_safe_text_edit_keeps_styles_tables_and_header():
    original = _sample_word()
    content = import_content(original, UPLOAD_ID)
    edited = content.replace("Revenue increased", "Sales increased")
    result = Document(io.BytesIO(render_edited_word(original, content, edited)))
    assert result.paragraphs[1].text == "Sales increased"
    assert result.paragraphs[1].runs[1].bold is True
    assert result.tables[0].cell(0, 0).text == "Table value"
    assert result.sections[0].header.paragraphs[0].text == "Confidential"


def test_cross_style_and_structure_edits_are_rejected():
    original = _sample_word()
    content = import_content(original, UPLOAD_ID)
    with pytest.raises(UnsafeWordEdit):
        render_edited_word(original, content, content.replace("Revenue increased", "Fell"))
    with pytest.raises(UnsafeWordEdit):
        render_edited_word(original, content, content + "\nNew paragraph")
    with pytest.raises(UnsafeWordEdit):
        render_edited_word(original, content, content.replace(UPLOAD_ID, "b" * 32 + ".docx"))


def test_word_import_edit_and_export_route_keeps_original_layout(tmp_path, monkeypatch):
    from core.database import Base
    from routes import document_routes
    from src.upload_handler import UploadHandler
    from src import word_preview

    engine = create_engine(f"sqlite:///{tmp_path / 'documents.db'}")
    Base.metadata.create_all(engine)
    monkeypatch.setattr(document_routes, "SessionLocal", sessionmaker(bind=engine))
    monkeypatch.setattr(document_routes, "get_current_user", lambda request: "alice")
    monkeypatch.setattr("src.auth_helpers.require_privilege", lambda request, privilege: "alice")
    app = FastAPI()
    app.include_router(document_routes.setup_document_routes(None, UploadHandler(str(tmp_path), str(tmp_path / 'uploads'))))
    original = _sample_word()
    previewed = []
    def fake_preview(word_bytes):
        previewed.append(word_bytes)
        return b"%PDF-1.4\n%%EOF"
    monkeypatch.setattr(word_preview, "render_word_pdf", fake_preview)

    with TestClient(app) as client:
        imported = client.post('/api/documents/import-docx', files={"file": ("Report.docx", original)} )
        assert imported.status_code == 200, imported.text
        document = imported.json()
        doc_id = document['id']
        assert client.get(f'/api/document/{doc_id}/export-docx').content == original
        initial_preview = client.get(f'/api/document/{doc_id}/render-docx')
        assert initial_preview.status_code == 200
        assert initial_preview.headers['content-type'] == 'application/pdf'
        assert initial_preview.headers['cache-control'] == 'private, no-store'
        assert previewed[-1] == original

        edited = document['current_content'].replace('Revenue increased', 'Sales increased')
        saved = client.put(f'/api/document/{doc_id}', json={"content": edited})
        assert saved.status_code == 200, saved.text
        exported = client.get(f'/api/document/{doc_id}/export-docx')
        assert exported.status_code == 200, exported.text
        result = Document(io.BytesIO(exported.content))
        assert result.paragraphs[1].text == 'Sales increased'
        assert result.paragraphs[1].runs[1].bold
        edited_preview = client.get(f'/api/document/{doc_id}/render-docx')
        assert edited_preview.status_code == 200
        assert previewed[-1] == exported.content

        unsafe = client.put(f'/api/document/{doc_id}', json={"content": edited + '\nNew paragraph'})
        assert unsafe.status_code == 422
        assert client.get(f'/api/document/{doc_id}/export-docx').content == exported.content


def test_upload_cleanup_preserves_document_sources(tmp_path, monkeypatch):
    from core import database
    from src.upload_handler import UploadHandler

    engine = create_engine(f"sqlite:///{tmp_path / 'cleanup.db'}")
    database.Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine)
    monkeypatch.setattr(database, 'SessionLocal', sessions)
    source_dir = tmp_path / 'uploads' / '2000' / '01' / '01'
    source_dir.mkdir(parents=True)
    source = source_dir / UPLOAD_ID
    source.write_bytes(_sample_word())
    expired = source_dir / ('b' * 32 + '.txt')
    expired.write_text('old attachment')
    with sessions() as db:
        db.add(database.Document(id='word', title='Word', language='docx',
                                 current_content=import_content(source.read_bytes(), UPLOAD_ID),
                                 owner='alice'))
        db.commit()
    handler = UploadHandler(str(tmp_path), str(tmp_path / 'uploads'))
    assert handler.cleanup_old_uploads() == 1
    assert source.exists()
    assert not expired.exists()


def test_chat_attachment_opens_as_word_document_in_library(tmp_path, monkeypatch):
    from core import database
    from src import database as database_alias
    from src.document_processor import build_user_content

    engine = create_engine(f"sqlite:///{tmp_path / 'chat.db'}")
    database.Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine)
    monkeypatch.setattr(database_alias, 'SessionLocal', sessions)
    monkeypatch.setattr(database_alias, 'Document', database.Document, raising=False)
    monkeypatch.setattr(database_alias, 'DocumentVersion', database.DocumentVersion, raising=False)
    monkeypatch.setattr(database_alias, 'Session', database.Session, raising=False)
    with sessions() as db:
        db.add(database.Session(id='chat', name='Chat', endpoint_url='local', model='test', owner='alice'))
        db.commit()

    source = tmp_path / UPLOAD_ID
    source.write_bytes(_sample_word())

    class Uploads:
        def resolve_upload(self, upload_id, owner=None):
            return {'path': str(source), 'name': 'Report.docx', 'mime': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'} if upload_id == UPLOAD_ID and owner == 'alice' else None

        def _inside_upload_dir(self, path):
            return path == str(source)

        def is_image_file(self, name, mime):
            return False

        def is_audio_file(self, name, mime):
            return False

        def is_document_file(self, name, mime):
            return True

    opened = []
    build_user_content('Read this', [UPLOAD_ID], str(tmp_path), Uploads(),
                       session_id='chat', auto_opened_docs=opened, owner='alice')
    assert len(opened) == 1
    assert opened[0]['title'] == 'Report'
    assert opened[0]['language'] == 'docx'
    assert opened[0]['content'].startswith(f'<!-- word_source upload_id="{UPLOAD_ID}" -->')
