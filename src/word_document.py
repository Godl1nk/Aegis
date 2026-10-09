"""Lossless-source Word imports with safe in-place text edits."""

from __future__ import annotations

import io
import re
import zipfile
from copy import deepcopy
from difflib import SequenceMatcher

from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from src.upload_handler import is_valid_upload_id


_SOURCE = re.compile(r'^<!-- word_source upload_id="([^"]+)" -->\n')


class UnsafeWordEdit(ValueError):
    """The requested change cannot preserve the imported Word layout."""


def source_upload_id(content: str) -> str | None:
    match = _SOURCE.match(content or "")
    if not match or not is_valid_upload_id(match.group(1)):
        return None
    return match.group(1)


def _paragraphs(document):
    return [Paragraph(node, document) for node in document.element.body.xpath('.//w:p')]


def import_content(data: bytes, upload_id: str) -> str:
    if not is_valid_upload_id(upload_id):
        raise ValueError("Invalid Word source ID")
    document = WordDocument(io.BytesIO(data))
    paragraphs = _paragraphs(document)
    return f'<!-- word_source upload_id="{upload_id}" -->\n' + '\n'.join(p.text for p in paragraphs)


def render_edited_word(source: bytes, original_content: str, edited_content: str) -> bytes:
    """Change paragraph text across plain runs while preserving OOXML objects."""
    source_id = source_upload_id(original_content)
    if not source_id or source_upload_id(edited_content) != source_id:
        raise UnsafeWordEdit("The Word source reference must stay intact")
    if edited_content == original_content:
        return source
    document = WordDocument(io.BytesIO(source))
    paragraphs = _paragraphs(document)
    original = original_content.split('\n', 1)[1]
    edited = edited_content.split('\n', 1)[1]
    original_lines = original.split('\n')
    edited_lines = edited.split('\n')
    if not paragraphs and not original:
        paragraphs = [document.add_paragraph()]
    if len(original_lines) != len(paragraphs):
        raise UnsafeWordEdit("The imported Word paragraphs could not be mapped")
    if [p.text for p in paragraphs] != original_lines:
        raise UnsafeWordEdit("The Word source no longer matches the imported text")
    changes = _paragraph_changes(original_lines, edited_lines)
    for action, old_start, old_end, new_start, new_end in changes:
        old_count, new_count = old_end - old_start, new_end - new_start
        if action == 'equal':
            continue
        if action == 'insert':
            _insert_word_paragraphs(paragraphs, old_start, edited_lines[new_start:new_end],
                                    at_document_end=old_start == len(paragraphs))
            continue
        if action == 'delete' or new_count == 0 or (old_count != new_count and old_count != 1):
            raise UnsafeWordEdit("The edit removes or ambiguously moves Word paragraphs")
        for offset in range(old_count):
            after = edited_lines[new_start + offset]
            _edit_word_paragraph(paragraphs[old_start + offset], original_lines[old_start + offset], after)
        if new_count > old_count:
            _insert_word_paragraphs(paragraphs, old_end, edited_lines[new_start + old_count:new_end])
    # The rendered paragraph order must match the editor before packaging.
    if [p.text for p in _paragraphs(document)] != edited_lines:
        raise UnsafeWordEdit("The Word paragraph positions could not be verified")
    # Saving through python-docx rewrites every package part. Only the main
    # document XML changed here; copy all other parts from the original file.
    main_part = str(document.part.partname).lstrip('/')
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as original_zip, zipfile.ZipFile(result, 'w') as edited_zip:
        members = original_zip.infolist()
        if sum(info.filename == main_part for info in members) != 1:
            raise UnsafeWordEdit("The Word main document part is missing or duplicated")
        if any(info.filename.startswith('_xmlsignatures/') for info in members):
            raise UnsafeWordEdit("Editing would invalidate the document's digital signature")
        edited_zip.comment = original_zip.comment
        for info in members:
            data = document.part.blob if info.filename == main_part else original_zip.read(info)
            edited_zip.writestr(info, data)
    edited_bytes = result.getvalue()
    if [p.text for p in _paragraphs(WordDocument(io.BytesIO(edited_bytes)))] != edited_lines:
        raise UnsafeWordEdit("The edited Word file could not be verified")
    return edited_bytes


def _paragraph_changes(original: list[str], edited: list[str]):
    # Pin unchanged boundaries before matching repeated template text. A global
    # longest-match diff can align a later blank section with an earlier one.
    start = 0
    while start < min(len(original), len(edited)) and original[start] == edited[start]:
        start += 1
    old_end, new_end = len(original), len(edited)
    while old_end > start and new_end > start and original[old_end - 1] == edited[new_end - 1]:
        old_end -= 1
        new_end -= 1
    if len(original) == len(edited):
        return [('replace', start, old_end, start, new_end)]
    return [(action, start + i, start + j, start + k, start + l)
            for action, i, j, k, l in SequenceMatcher(
                None, original[start:old_end], edited[start:new_end], autojunk=False
            ).get_opcodes()]


def _edit_word_paragraph(paragraph, before: str, after: str) -> None:
    if before == after:
        return
    start = 0
    while start < min(len(before), len(after)) and before[start] == after[start]:
        start += 1
    suffix = 0
    while (suffix < len(before) - start and suffix < len(after) - start
           and before[-suffix - 1] == after[-suffix - 1]):
        suffix += 1
    end = len(before) - suffix
    replacement_end = len(after) - suffix
    runs = paragraph.runs
    if ''.join(run.text for run in runs) != before:
        raise UnsafeWordEdit("This paragraph contains text outside ordinary Word runs")
    if not runs and not before and all(
        node.tag == qn('w:pPr') for node in paragraph._p
    ):
        runs = [paragraph.add_run('')]
        runs[0]._element.add_t('')
    spans = []
    offset = 0
    for run in runs:
        spans.append((run, offset, offset + len(run.text)))
        offset += len(run.text)
    if start == end:
        touched = [span for span in spans if span[1] <= start <= span[2]
                   and _plain_text_run(span[0])][:1]
    else:
        touched = [span for span in spans if span[1] < end and span[2] > start]
    if not touched:
        raise UnsafeWordEdit("The edit could not be located in Word text")
    for run, _, _ in touched:
        if not _plain_text_run(run):
            raise UnsafeWordEdit("The edit touches a Word field or object")
    first_run, first_start, _ = touched[0]
    last_run, last_start, last_end = touched[-1]
    replacement = after[start:replacement_end]
    if first_run is last_run:
        _set_run_text(first_run, before[first_start:start] + replacement + before[end:last_end])
    else:
        _set_run_text(first_run, before[first_start:start] + replacement)
        for run, _, _ in touched[1:-1]:
            _set_run_text(run, '')
        _set_run_text(last_run, before[end:last_end])
    if paragraph.text != after:
        raise UnsafeWordEdit("The edited Word text could not be verified")


def _insert_word_paragraphs(paragraphs, index: int, lines: list[str],
                            at_document_end: bool = False) -> None:
    anchor = paragraphs[index - 1] if index else paragraphs[0]
    if at_document_end and anchor._p.getparent().tag != qn('w:body'):
        for line in lines:
            anchor._parent.add_paragraph(line)
        return
    for line in lines:
        node = OxmlElement('w:p')
        if anchor._p.pPr is not None:
            properties = deepcopy(anchor._p.pPr)
            for tag in ('w:sectPr', 'w:pageBreakBefore'):
                for unsafe in properties.xpath(f'./{tag}'):
                    properties.remove(unsafe)
            node.append(properties)
        if index:
            anchor._p.addnext(node)
        else:
            anchor._p.addprevious(node)
        new_paragraph = Paragraph(node, anchor._parent)
        new_paragraph.add_run(line)
        if index:
            anchor = new_paragraph


def _set_run_text(run, value: str) -> None:
    """Keep run properties and XML nodes; change only their text payload."""
    text_nodes = [node for node in run._element if node.tag == qn('w:t')]
    text_nodes[0].text = value
    for node in text_nodes[1:]:
        node.text = ''


def _plain_text_run(run) -> bool:
    children = list(run._element)
    return any(node.tag == qn('w:t') for node in children) and all(
        node.tag in (qn('w:rPr'), qn('w:t')) for node in children
    )


def render_document_edit(db, doc, edited_content: str, upload_handler, owner: str | None, auth_manager=None) -> bytes:
    """Replay saved edits in order, preserving paragraph identity between versions."""
    from core.database import DocumentVersion

    versions = db.query(DocumentVersion).filter(
        DocumentVersion.document_id == doc.id,
    ).order_by(DocumentVersion.version_number).all()
    if not versions or versions[0].version_number != 1:
        raise UnsafeWordEdit("The original Word version is unavailable")
    original = versions[0]
    upload_id = source_upload_id(original.content)
    if not upload_id or source_upload_id(edited_content) != upload_id:
        raise UnsafeWordEdit("The Word source reference must stay intact")
    resolved = upload_handler.resolve_upload(upload_id, owner=owner, auth_manager=auth_manager)
    if not resolved:
        raise UnsafeWordEdit("The original Word file is unavailable")
    with open(resolved['path'], 'rb') as file:
        result = file.read()
    content = original.content
    rendered = {content: result}
    # Restores and historical exports stop at the existing snapshot rather than
    # interpreting a restore as deletion of subsequently inserted paragraphs.
    target = next((version for version in versions if version.content == edited_content), None)
    current = next((version for version in reversed(versions)
                    if version.content == doc.current_content), None)
    stop = target or current
    if stop is None:
        raise UnsafeWordEdit("The current Word version is unavailable")
    for version in versions[1:]:
        if version.version_number > stop.version_number:
            break
        if version.content in rendered:
            result = rendered[version.content]
        else:
            result = render_edited_word(result, content, version.content)
            rendered[version.content] = result
        content = version.content
    return render_edited_word(result, content, edited_content)
